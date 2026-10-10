import io
import zipfile

import pytest

from app.services.ingestion.extractors.docx import extract_docx
from app.services.ingestion.extractors.pdf import extract_pdf
from app.services.ingestion.extractors.txt import decode_text, extract_txt
from app.services.ingestion.normalize import assemble, normalize_text
from app.services.ingestion.types import BlockKind, ExtractionError
from tests.factories import blank_pdf, encrypt_pdf, make_docx, make_pdf, write

LEAVE = (
    "Employees are entitled to fourteen days of annual leave per calendar year. "
    "Leave must be approved by the line manager at least two weeks in advance."
)


# ---------------------------------------------------------------- normalisation
@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("ofﬁce", "office"),  # ligature folded by NFKC
        ("infor-\nmation", "information"),  # line-break hyphenation repaired
        ("Kuala-\nLumpur", "Kuala- Lumpur"),  # capitalised compound left alone
        ("a­b", "ab"),  # soft hyphen removed
        ("line one\nline two", "line one line two"),  # wrapped lines joined
        ("x\x00y\x07z", "x y z"),  # control chars removed
        ("  many    spaces here ", "many spaces here"),
    ],
)
def test_normalize_text(raw, expected):
    assert normalize_text(raw) == expected


# ---------------------------------------------------------------------- PDF
def test_pdf_keeps_page_numbers_headings_and_metadata(tmp_path):
    path = write(
        tmp_path,
        "h.pdf",
        make_pdf(
            [["Annual Leave Policy", LEAVE], ["Claims are reimbursed within 30 days."]],
            outline={1: "Leave", 2: "Claims"},
            title="HR Handbook",
        ),
    )
    doc = extract_pdf(path, max_pages=10)
    assert doc.page_count == 2
    assert doc.metadata["title"] == "HR Handbook"
    by_page = {b.page: b for b in doc.blocks}
    assert "fourteen days" in " ".join(b.text for b in doc.blocks if b.page == 1)
    assert by_page[2].text == "Claims are reimbursed within 30 days."
    assert by_page[1].heading == "Leave" and by_page[2].heading == "Claims"
    # The title line is split off from the paragraph below it.
    assert doc.blocks[0].text == "Annual Leave Policy"


def test_pdf_without_outline_has_no_headings(tmp_path):
    path = write(tmp_path, "p.pdf", make_pdf([[LEAVE]]))
    assert all(b.heading is None for b in extract_pdf(path, max_pages=10).blocks)


def test_pdf_with_user_password_is_rejected(tmp_path):
    path = write(tmp_path, "e.pdf", encrypt_pdf(make_pdf([[LEAVE]]), "s3cret"))
    with pytest.raises(ExtractionError) as err:
        extract_pdf(path, max_pages=10)
    assert err.value.code == "encrypted"


def test_pdf_with_empty_user_password_is_read(tmp_path):
    path = write(tmp_path, "e.pdf", encrypt_pdf(make_pdf([[LEAVE]]), ""))
    assert "fourteen" in extract_pdf(path, max_pages=10).blocks[0].text


def test_scanned_pdf_without_text_fails_clearly(tmp_path):
    path = write(tmp_path, "scan.pdf", blank_pdf(3))
    with pytest.raises(ExtractionError) as err:
        extract_pdf(path, max_pages=10)
    assert err.value.code == "no_text"
    assert "OCR" in err.value.message


def test_pdf_page_limit(tmp_path):
    path = write(tmp_path, "big.pdf", make_pdf([[LEAVE]] * 3))
    with pytest.raises(ExtractionError) as err:
        extract_pdf(path, max_pages=2)
    assert err.value.code == "too_many_pages"


def test_corrupt_pdf(tmp_path):
    path = write(tmp_path, "bad.pdf", b"%PDF-1.7\n this is not really a pdf")
    with pytest.raises(ExtractionError) as err:
        extract_pdf(path, max_pages=10)
    assert err.value.code in {"corrupt_file", "no_text", "empty_document"}


# --------------------------------------------------------------------- DOCX
def test_docx_headings_tables_and_metadata(tmp_path):
    data = make_docx(
        [
            ("title", "HR Handbook"),
            ("h1", "Leave"),
            ("p", LEAVE),
            ("h2", "Medical"),
            ("p", "Bring a medical certificate."),
        ],
        title="HR Handbook",
        author="Maksuda",
        table=[["Grade", "Days"], ["A", "14"], ["B", "18"]],
    )
    doc = extract_docx(write(tmp_path, "h.docx", data), max_uncompressed_bytes=10**8)
    kinds = [(b.kind, b.text) for b in doc.blocks]
    assert (BlockKind.HEADING, "Leave") in kinds
    assert (BlockKind.TABLE, "Grade | Days\nA | 14\nB | 18") in kinds
    para = next(b for b in doc.blocks if b.text.startswith("Bring"))
    assert para.heading == "Medical" and para.page is None
    assert doc.metadata["title"] == "HR Handbook" and doc.metadata["author"] == "Maksuda"
    assert doc.page_count is None


def test_docx_merged_cells_appear_once(tmp_path):
    import docx as python_docx

    d = python_docx.Document()
    t = d.add_table(rows=2, cols=3)
    t.cell(0, 0).merge(t.cell(0, 2)).text = "Merged header"
    for i, v in enumerate(["a", "b", "c"]):
        t.cell(1, i).text = v
    buf = io.BytesIO()
    d.save(buf)
    doc = extract_docx(write(tmp_path, "m.docx", buf.getvalue()), max_uncompressed_bytes=10**8)
    assert doc.blocks[0].text == "Merged header\na | b | c"


def test_docx_zip_bomb_rejected_before_parsing(tmp_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<x/>")
        zf.writestr("word/document.xml", "<x/>")
        zf.writestr("word/padding.bin", b"\0" * 5_000_000)  # compresses to a few KB
    path = write(tmp_path, "bomb.docx", buf.getvalue())
    with pytest.raises(ExtractionError) as err:
        extract_docx(path, max_uncompressed_bytes=1_000_000)
    assert err.value.code == "archive_too_large"


def test_empty_docx(tmp_path):
    path = write(tmp_path, "e.docx", make_docx([]))
    with pytest.raises(ExtractionError) as err:
        extract_docx(path, max_uncompressed_bytes=10**8)
    assert err.value.code == "no_text"


# ---------------------------------------------------------------------- TXT
def test_txt_markdown_headings_and_paragraphs(tmp_path):
    text = "# Leave\nAnnual leave is 14 days.\nIt is pro-rated.\n\n## Claims\n\nPaid monthly.\n"
    doc = extract_txt(write(tmp_path, "n.txt", text.encode()))
    assert [(b.kind.value, b.text, b.heading) for b in doc.blocks] == [
        ("heading", "Leave", "Leave"),
        ("paragraph", "Annual leave is 14 days. It is pro-rated.", "Leave"),
        ("heading", "Claims", "Claims"),
        ("paragraph", "Paid monthly.", "Claims"),
    ]
    assert doc.metadata["encoding"] == "utf-8" and not doc.warnings


@pytest.mark.parametrize(
    ("raw", "encoding"),
    [
        ("Café résumé".encode(), "utf-8"),
        ("﻿BOM text".encode(), "utf-8"),
        ("Café 50€".encode("cp1252"), "cp1252"),
    ],
)
def test_txt_encoding_detection(raw, encoding):
    text, detected = decode_text(raw)
    assert detected == encoding
    assert not text.startswith("﻿")


def test_non_utf8_txt_records_warning(tmp_path):
    doc = extract_txt(write(tmp_path, "w.txt", "Café".encode("cp1252")))
    assert doc.blocks[0].text == "Café"
    assert doc.warnings == ["decoded as cp1252, not UTF-8"]


def test_whitespace_only_txt(tmp_path):
    with pytest.raises(ExtractionError):
        extract_txt(write(tmp_path, "w.txt", b"  \n\n \n"))


# ----------------------------------------------------------------- assembly
def test_assemble_offsets_point_at_block_text(tmp_path):
    path = write(tmp_path, "h.pdf", make_pdf([["First para here."], ["Second page text."]]))
    extracted = extract_pdf(path, max_pages=10)
    assembled = assemble(extracted.blocks)
    for block, span in zip(extracted.blocks, assembled.spans, strict=True):
        assert assembled.text[span.start : span.end] == block.text
        assert span.page == block.page
