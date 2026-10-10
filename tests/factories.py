"""Build real PDF / DOCX / TXT files for tests, so extraction is tested end to end."""

import io
from pathlib import Path

import docx
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


def make_pdf(
    pages: list[list[str]],
    *,
    outline: dict[int, str] | None = None,
    title: str | None = None,
) -> bytes:
    """pages: list of pages, each a list of paragraphs. Paragraphs are separated by a gap.

    outline: {1-based page number: bookmark title}.
    """
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    if title:
        c.setTitle(title)
    for page_no, paragraphs in enumerate(pages, start=1):
        if outline and page_no in outline:
            key = f"p{page_no}"
            c.bookmarkPage(key)
            c.addOutlineEntry(outline[page_no], key, level=0)
        y = 800
        for para in paragraphs:
            for line in _wrap(para, 85):
                c.drawString(50, y, line)
                y -= 14
            y -= 14  # blank line between paragraphs
        c.showPage()
    c.save()
    return buf.getvalue()


def encrypt_pdf(data: bytes, password: str) -> bytes:
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(data)))
    writer.encrypt(user_password=password, owner_password=password, algorithm="RC4-128")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def blank_pdf(pages: int = 1) -> bytes:
    """A PDF with no text at all — stands in for a scanned document."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for _ in range(pages):
        c.rect(100, 100, 200, 200, fill=1)
        c.showPage()
    c.save()
    return buf.getvalue()


def make_docx(
    items: list[tuple[str, str]],
    *,
    title: str | None = None,
    author: str | None = None,
    table: list[list[str]] | None = None,
) -> bytes:
    """items: (kind, text) where kind is 'h1', 'h2', 'p' or 'title'."""
    document = docx.Document()
    if title:
        document.core_properties.title = title
    if author:
        document.core_properties.author = author
    for kind, text in items:
        if kind == "title":
            document.add_heading(text, level=0)
        elif kind.startswith("h"):
            document.add_heading(text, level=int(kind[1:]))
        else:
            document.add_paragraph(text)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for col, value in enumerate(row):
                t.cell(r, col).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def write(tmp_path: Path, name: str, data: bytes) -> Path:
    path = tmp_path / name
    path.write_bytes(data)
    return path


def _wrap(text: str, width: int) -> list[str]:
    lines, current = [], ""
    for word in text.split():
        if current and len(current) + 1 + len(word) > width:
            lines.append(current)
            current = word
        else:
            current = f"{current} {word}".strip()
    if current:
        lines.append(current)
    return lines
